"""Extract text from uploaded PDF/DOCX files and parse it into profile suggestions.

Extraction is best-effort. Anything that cannot be parsed with confidence is
marked as unconfirmed and surfaced for user review; nothing is written to the
profile automatically.
"""
from __future__ import annotations

import io
import logging
import re
from datetime import date

from app.core.errors import ValidationAppError
from app.core.text import KNOWN_SKILLS, clean_whitespace, extract_skills

logger = logging.getLogger("careerpilot.extract")

SECTION_ALIASES = {
    "summary": {"summary", "profile", "professional summary", "about", "objective", "career summary"},
    "experience": {"experience", "work experience", "professional experience", "employment",
                   "employment history", "work history"},
    "education": {"education", "academic background", "academics"},
    "skills": {"skills", "technical skills", "core skills", "technologies", "tech stack",
               "competencies"},
    "projects": {"projects", "selected projects", "personal projects", "portfolio"},
    "certifications": {"certifications", "certificates", "licenses", "credentials"},
    "achievements": {"achievements", "awards", "honors", "recognition"},
    "contact": {"contact", "contact details", "links"},
}

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d[\d\s().-]{7,}\d)")
URL_RE = re.compile(r"https?://[^\s)>\]]+")
DATE_RE = re.compile(
    r"((?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*\d{4}|\d{1,2}/\d{4}|\d{4})",
    re.IGNORECASE,
)
BULLET_RE = re.compile(r"^\s*[•●▪◦\-\*\u2013\u2022]\s+")
COMPANY_HINT_RE = re.compile(r"\b(?:at|@)\s+([A-Z][\w&.\- ]{1,60})")


def extract_text(data: bytes, fmt: str) -> str:
    if fmt == "pdf":
        return _extract_pdf(data)
    if fmt == "docx":
        return _extract_docx(data)
    raise ValidationAppError("Unsupported document format")


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is pinned
        raise ValidationAppError("PDF extraction is unavailable on this deployment") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > 20:
            raise ValidationAppError("Please upload a document with 20 pages or fewer.")
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ValidationAppError("Password-protected PDFs are not supported.") from exc
        parts = []
        for page in reader.pages[:20]:
            parts.append(page.extract_text() or "")
        text = "\n".join(parts)
    except ValidationAppError:
        raise
    except Exception as exc:
        logger.warning("pdf_extract_failed error=%s", type(exc).__name__)
        raise ValidationAppError("That PDF could not be read. Try a DOCX or paste the text.") from exc
    if len(clean_whitespace(text)) < 30:
        raise ValidationAppError(
            "No selectable text was found in this PDF (it may be a scan). Upload a DOCX or paste your details."
        )
    return text[:120_000]


def _extract_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ValidationAppError("DOCX extraction is unavailable on this deployment") from exc
    try:
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
    except Exception as exc:
        logger.warning("docx_extract_failed error=%s", type(exc).__name__)
        raise ValidationAppError("That DOCX could not be read.") from exc
    if len(clean_whitespace(text)) < 30:
        raise ValidationAppError("That DOCX appears to be empty.")
    return text[:120_000]


def _split_sections(text: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {}
    current = "header"
    sections[current] = []
    lookup = {alias: key for key, aliases in SECTION_ALIASES.items() for alias in aliases}
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        normalized = re.sub(r"[^a-z ]", "", line.lower()).strip()
        if normalized in lookup and len(line) <= 60:
            current = lookup[normalized]
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return sections


def _parse_date_range(line: str) -> tuple[date | None, date | None, bool]:
    matches = DATE_RE.findall(line or "")
    if not matches:
        return None, None, "present" in (line or "").lower()
    parsed: list[date] = []
    for match in matches[:4]:
        try:
            if re.fullmatch(r"\d{4}", match):
                parsed.append(date(int(match), 1, 1))
            elif "/" in match:
                month, year = match.split("/")
                parsed.append(date(int(year), max(1, min(12, int(month))), 1))
            else:
                from datetime import datetime

                dt = datetime.strptime(match.replace(".", ""), "%b %Y")
                parsed.append(date(dt.year, dt.month, 1))
        except (ValueError, TypeError):
            continue
    start = parsed[0] if parsed else None
    end = parsed[1] if len(parsed) > 1 else None
    present = bool(re.search(r"\bpresent\b|\bcurrent\b", line, re.IGNORECASE))
    if present:
        end = None
    return start, end, present


def parse_resume_text(text: str) -> dict:
    """Heuristic structure extraction. Every field is marked unconfirmed."""
    sections = _split_sections(text)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    header_lines = sections.get("header", [])[:8]
    full_header = " ".join(header_lines)

    email_match = EMAIL_RE.search(text)
    phone_match = PHONE_RE.search(full_header)
    links = URL_RE.findall(text)
    github = next((link for link in links if "github.com" in link.lower()), "")
    linkedin = next((link for link in links if "linkedin.com" in link.lower()), "")
    portfolio = next(
        (link for link in links if "github.com" not in link.lower() and "linkedin.com" not in link.lower()),
        "",
    )

    name = ""
    for line in header_lines:
        if EMAIL_RE.search(line) or URL_RE.search(line) or PHONE_RE.search(line):
            continue
        if len(line.split()) <= 6 and not re.search(r"\d{3,}", line):
            name = line[:120]
            break

    headline = ""
    for line in header_lines:
        if line != name and len(line.split()) <= 12 and not EMAIL_RE.search(line) \
                and not URL_RE.search(line) and not PHONE_RE.search(line):
            headline = line[:200]
            break

    experience: list[dict] = []
    current: dict | None = None
    for line in sections.get("experience", []):
        is_bullet = bool(BULLET_RE.match(line))
        starts_with_date = bool(DATE_RE.search(line)) and len(line) < 160
        has_company_hint = bool(COMPANY_HINT_RE.search(line)) or " | " in line
        looks_like_title = (
            not is_bullet and len(line) < 140 and
            (starts_with_date or has_company_hint) and
            (re.search(r"\b(engineer|developer|manager|analyst|designer|scientist|consultant|"
                       r"intern|lead|architect|director|specialist|administrator|associate|"
                       r"founder|researcher|officer|coordinator)\b", line, re.IGNORECASE) is not None)
        )
        if looks_like_title and (current is None or current.get("description")):
            start, end, present = _parse_date_range(line)
            company = ""
            company_match = COMPANY_HINT_RE.search(line)
            if company_match:
                company = company_match.group(1).strip()
            title_part = re.split(r"\s+[|@•·–—-]\s+", line)[0]
            title_part = re.sub(r"\(?(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*\d{4}.*$",
                                "", title_part, flags=re.IGNORECASE).strip(" -|,")
            current = {
                "title": clean_whitespace(title_part)[:200],
                "employer": clean_whitespace(company)[:200],
                "start_date": start.isoformat() if start else None,
                "end_date": end.isoformat() if end else None,
                "is_current": present,
                "description": "",
                "achievements": [],
                "provenance": "document_extracted",
                "verified": False,
            }
            experience.append(current)
        elif current is not None:
            if is_bullet:
                current["achievements"].append(BULLET_RE.sub("", line)[:400])
            else:
                current["description"] = (current["description"] + " " + line).strip()[:2000]
    for item in experience:
        item["achievements"] = item["achievements"][:12]

    education: list[dict] = []
    for line in sections.get("education", [])[:12]:
        if BULLET_RE.match(line):
            continue
        if len(line) < 4:
            continue
        start, end, _present = _parse_date_range(line)
        degree_match = re.search(
            r"\b(bachelor|b\.?sc|b\.?a|master|m\.?sc|mba|phd|doctorate|diploma|certificate|"
            r"associate|hnd|ond)\b", line, re.IGNORECASE)
        institution = re.split(r"\s+[|•·,–—-]\s+|\d{4}", line)[0].strip(" ,-|")[:250]
        education.append({
            "institution": institution or line[:250],
            "qualification": " ".join(line.split())[:200] if degree_match else "",
            "field_of_study": "",
            "start_date": start.isoformat() if start else None,
            "end_date": end.isoformat() if end else None,
            "status": "completed",
            "provenance": "document_extracted",
            "verified": False,
        })

    skills_text = " ".join(sections.get("skills", []))
    skills_found = extract_skills(skills_text) or extract_skills(text[:20_000])
    skills = [
        {"name": skill[:120], "category": _guess_category(skill), "proficiency": "unspecified",
         "verified": False, "provenance": "document_extracted", "evidence": ""}
        for skill in skills_found[:40]
    ]

    projects: list[dict] = []
    for line in sections.get("projects", [])[:10]:
        if BULLET_RE.match(line):
            continue
        parts = re.split(r"\s+[|–—]\s+", line, maxsplit=1)
        if len(line) > 3:
            projects.append({
                "name": parts[0][:200],
                "description": (parts[1] if len(parts) > 1 else "")[:2000],
                "technologies": extract_skills(line)[:12],
                "provenance": "document_extracted",
                "verified": False,
            })

    certifications = []
    for line in sections.get("certifications", [])[:10]:
        if len(line) > 2:
            certifications.append({
                "name": BULLET_RE.sub("", line)[:250],
                "issuer": "",
                "provenance": "document_extracted",
                "verified": False,
            })

    summary = " ".join(sections.get("summary", []))[:2000]

    warnings: list[str] = []
    if not name:
        warnings.append("A name could not be identified. Enter it manually.")
    if not experience:
        warnings.append("No work experience was detected with confidence. Add positions manually.")
    if not email_match:
        warnings.append("No email address was found in the document.")
    for item in experience:
        if not item["employer"]:
            warnings.append(f"Employer missing for “{item['title'] or 'a position'}”. Confirm before saving.")
            break

    return {
        "profile": {
            "preferred_name": name,
            "professional_headline": headline,
            "summary": summary,
            "email": email_match.group(0) if email_match else "",
            "phone": clean_whitespace(phone_match.group(0)) if phone_match else "",
            "github_url": github[:500],
            "linkedin_url": linkedin[:500],
            "portfolio_url": portfolio[:500],
            "provenance": "document_extracted",
        },
        "experiences": experience,
        "education": education,
        "skills": skills,
        "projects": projects,
        "certifications": certifications,
        "warnings": warnings[:10],
        "suggestions": _suggestions(skills, experience, education, summary),
        "text_preview": clean_whitespace(text)[:1200],
        "confidence": {
            "name": "high" if name else "low",
            "experience": "medium" if experience else "low",
            "skills": "high" if len(skills) >= 3 else "low",
            "education": "medium" if education else "low",
        },
        "source": "document_extracted",
        "requires_review": True,
    }


def _guess_category(skill: str) -> str:
    if skill in {"python", "javascript", "typescript", "java", "go", "rust", "c++", "c#",
                 "ruby", "php", "swift", "kotlin", "sql", "html", "css", "dart"}:
        return "Languages"
    if skill in {"react", "next.js", "vue", "angular", "svelte", "tailwind", "redux", "node.js"}:
        return "Frontend"
    if skill in {"django", "flask", "fastapi", "spring", "express", "rails", "laravel", ".net",
                 "postgresql", "mysql", "mongodb", "redis", "graphql", "rest"}:
        return "Backend & Data"
    if skill in {"aws", "azure", "google cloud", "docker", "kubernetes", "terraform", "ci/cd",
                 "github actions", "jenkins", "linux", "nginx"}:
        return "Cloud & DevOps"
    if skill in {"machine learning", "deep learning", "tensorflow", "pytorch", "pandas", "nlp",
                 "llms", "data analysis", "spark"}:
        return "AI & Analytics"
    if skill in {"agile", "scrum", "leadership", "communication", "project management", "jira"}:
        return "Ways of Working"
    return "General"


def _suggestions(skills: list[dict], experience: list[dict], education: list[dict],
                 summary: str) -> list[str]:
    suggestions = []
    if not summary:
        suggestions.append("Add a short professional summary (2–3 sentences) to introduce your focus.")
    missing_metrics = [e for e in experience if not any(ch.isdigit() for ch in " ".join(e.get("achievements", [])))]
    if missing_metrics:
        suggestions.append("Add measurable outcomes to experience bullets where you can verify them.")
    if len(skills) < 5:
        suggestions.append("Add more skills with evidence of where you used them.")
    if not education:
        suggestions.append("Add education or relevant training if it strengthens your profile.")
    return suggestions[:5]
