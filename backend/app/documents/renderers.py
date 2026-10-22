"""DOCX, PDF, text and LaTeX renderers with output validation."""
from __future__ import annotations

import io
import logging
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from xml.sax.saxutils import escape

from app.core.text import strip_control_characters

logger = logging.getLogger("careerpilot.documents.render")

PAGE_WIDTHS = {"a4": 595.27, "letter": 612.0}
DEJAVU_DIRS = [
    Path("/usr/share/fonts/truetype/dejavu"),
    Path("/usr/share/fonts/dejavu"),
    Path("/usr/local/share/fonts/dejavu"),
]


class RenderError(Exception):
    pass


# --------------------------------------------------------------------------- text


def render_text(content: dict, doc_type: str) -> bytes:
    from app.documents.content import cover_letter_plain_text, cv_plain_text

    text = cover_letter_plain_text(content) if doc_type == "cover_letter" else cv_plain_text(content)
    if len(text.strip()) < 20:
        raise RenderError("The document content is too short to export.")
    return text.encode("utf-8")


# --------------------------------------------------------------------------- DOCX


def render_docx(content: dict, doc_type: str) -> bytes:
    try:
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
        from docx.shared import Pt, RGBColor, Inches
    except ImportError as exc:  # pragma: no cover
        raise RenderError("DOCX generation is unavailable on this deployment.") from exc

    document = docx.Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.05
    try:
        normal.element.rPr.rFonts.set(docx.oxml.ns.qn("w:eastAsia"), "Calibri")
    except Exception:  # pragma: no cover
        pass

    header = content.get("header", {})
    name = strip_control_characters(header.get("name", "") or "")
    title_paragraph = document.add_paragraph()
    run = title_paragraph.add_run(name)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x11, 0x1A, 0x13)
    title_paragraph.paragraph_format.space_after = Pt(0)

    if header.get("headline"):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(strip_control_characters(header["headline"]))
        run.italic = True
        run.font.size = Pt(10.5)
        paragraph.paragraph_format.space_after = Pt(2)

    contact_bits = [header.get("email"), header.get("phone"), header.get("location")]
    contact = "  |  ".join(bits for bits in contact_bits if bits)
    if contact:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(strip_control_characters(contact))
        run.font.size = Pt(9.5)
        paragraph.paragraph_format.space_after = Pt(1)
    links = header.get("links") or []
    if links:
        paragraph = document.add_paragraph()
        run = paragraph.add_run("  |  ".join(f"{link.get('label')}: {link.get('url')}" for link in links))
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x44, 0x38)
        paragraph.paragraph_format.space_after = Pt(4)

    if doc_type == "cover_letter":
        _docx_cover_letter(document, content)
    else:
        _docx_cv(document, content)

    buffer = io.BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    ok, message = validate_docx(data)
    if not ok:
        raise RenderError(message)
    return data


def _section_heading(document, text: str) -> None:
    from docx.shared import Pt, RGBColor

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(strip_control_characters(text).upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x14, 0x2A, 0x1A)


def _docx_cv(document, content: dict) -> None:
    from docx.shared import Pt

    if content.get("summary"):
        _section_heading(document, "Summary")
        document.add_paragraph(strip_control_characters(content["summary"]))
    for section in content.get("sections", []):
        title = section.get("title", "")
        _section_heading(document, title)
        if section.get("type") == "skills":
            for group in section.get("groups", []):
                paragraph = document.add_paragraph()
                run = paragraph.add_run(f"{group.get('category', 'Skills')}: ")
                run.bold = True
                paragraph.add_run(", ".join(strip_control_characters(item) for item in group.get("items", [])))
                paragraph.paragraph_format.space_after = Pt(1)
            continue
        for item in section.get("items", []):
            heading = document.add_paragraph()
            heading.paragraph_format.space_before = Pt(4)
            heading.paragraph_format.space_after = Pt(0)
            title_run = heading.add_run(strip_control_characters(item.get("title", "")))
            title_run.bold = True
            subtitle = item.get("subtitle") or ""
            if subtitle:
                heading.add_run(" — " + strip_control_characters(subtitle))
            if item.get("dates"):
                heading.add_run(f"  ({strip_control_characters(item['dates'])})")
            if item.get("location"):
                location_paragraph = document.add_paragraph()
                run = location_paragraph.add_run(strip_control_characters(item["location"]))
                run.italic = True
                run.font.size = Pt(9.5)
                location_paragraph.paragraph_format.space_after = Pt(1)
            for bullet in item.get("bullets", []) or []:
                text = strip_control_characters(str(bullet))
                if not text:
                    continue
                paragraph = document.add_paragraph(text, style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.left_indent = Pt(16)
            for link in item.get("links", []) or []:
                paragraph = document.add_paragraph(f"Link: {strip_control_characters(link)}", style="List Bullet")
                paragraph.paragraph_format.left_indent = Pt(16)


def _docx_cover_letter(document, content: dict) -> None:
    from docx.shared import Pt

    recipient = content.get("recipient", {})
    document.add_paragraph("")
    if content.get("date"):
        document.add_paragraph(strip_control_characters(content["date"]))
    if recipient.get("company"):
        p = document.add_paragraph(strip_control_characters(recipient["company"]))
        p.paragraph_format.space_after = Pt(0)
    if recipient.get("role"):
        document.add_paragraph(f"Re: {strip_control_characters(recipient['role'])}")
    document.add_paragraph(strip_control_characters(content.get("salutation", "Dear Hiring Manager,")))
    for paragraph_text in content.get("paragraphs", []):
        paragraph = document.add_paragraph(strip_control_characters(paragraph_text))
        paragraph.paragraph_format.space_after = Pt(8)
    document.add_paragraph("")
    document.add_paragraph(strip_control_characters(content.get("closing", "Sincerely,")))
    signature = document.add_paragraph(strip_control_characters(content.get("signature", "")))
    if signature.runs:
        signature.runs[0].bold = True


# --------------------------------------------------------------------------- PDF


def _font_paths() -> tuple[Path, Path, Path, Path] | None:
    for directory in DEJAVU_DIRS:
        regular = directory / "DejaVuSans.ttf"
        bold = directory / "DejaVuSans-Bold.ttf"
        italic = directory / "DejaVuSans-Oblique.ttf"
        bold_italic = directory / "DejaVuSans-BoldOblique.ttf"
        if regular.exists() and bold.exists():
            return regular, bold, italic if italic.exists() else regular, \
                bold_italic if bold_italic.exists() else bold
    bundled = Path(__file__).resolve().parent.parent / "assets" / "fonts"
    regular = bundled / "DejaVuSans.ttf"
    if regular.exists():
        return regular, bundled / "DejaVuSans-Bold.ttf", regular, regular
    return None


def _pdf_fonts() -> tuple[str, str, str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    paths = _font_paths()
    if not paths:
        return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Helvetica-BoldOblique"
    regular, bold, italic, bold_italic = paths
    try:
        pdfmetrics.registerFont(TTFont("CPRegular", str(regular)))
        pdfmetrics.registerFont(TTFont("CPBold", str(bold)))
        pdfmetrics.registerFont(TTFont("CPItalic", str(italic)))
        pdfmetrics.registerFont(TTFont("CPBoldItalic", str(bold_italic)))
        return "CPRegular", "CPBold", "CPItalic", "CPBoldItalic"
    except Exception:  # pragma: no cover - font issues must not break export
        logger.warning("pdf_font_registration_failed")
        return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Helvetica-BoldOblique"


def _pdf_safe(text: str, unicode_fonts: bool) -> str:
    text = strip_control_characters(str(text or ""))
    if unicode_fonts:
        return text
    replacements = {"—": "-", "–": "-", "•": "-", "’": "'", "‘": "'", "“": '"', "”": '"',
                    "…": "...", "→": "->", "≈": "~"}
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def _pdf_paragraph(text: str, style, unicode_fonts: bool):
    from reportlab.platypus import Paragraph

    escaped = escape(_pdf_safe(text, unicode_fonts))
    escaped = re.sub(r"(https?://[^\s<]+)", lambda m: _break_long_url(m.group(1)), escaped)
    return Paragraph(escaped, style)


def _break_long_url(url: str) -> str:
    if len(url) <= 70:
        return url
    chunks = [url[i:i + 60] for i in range(0, len(url), 60)]
    return "<br/>".join(chunks)


def _numbered_canvas_class():
    """Build a canvas subclass that can print accurate 'Page x of y' footers."""
    from reportlab.pdfgen import canvas

    class NumberedCanvas(canvas.Canvas):
        def __init__(self, *a, **kw):
            canvas.Canvas.__init__(self, *a, **kw)
            self._saved_page_states = []

        def showPage(self):  # noqa: N802 - reportlab API
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._saved_page_states)
            for state in self._saved_page_states:
                self.__dict__.update(state)
                self._draw_footer(total)
                canvas.Canvas.showPage(self)
            canvas.Canvas.save(self)

        def _draw_footer(self, total: int):
            self.setFont("Helvetica", 7.5)
            self.setFillColorRGB(0.45, 0.5, 0.46)
            self.drawRightString(self._pagesize[0] - 40, 28, f"Page {self._pageNumber} of {total}")

    return NumberedCanvas


def render_pdf(content: dict, doc_type: str, *, page_size: str = "a4") -> bytes:
    try:
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4, LETTER
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            KeepTogether,
            ListFlowable,
            ListItem,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
        )
    except ImportError as exc:  # pragma: no cover
        raise RenderError("PDF generation is unavailable on this deployment.") from exc

    fonts = _pdf_fonts()
    regular, bold, italic, _bold_italic = fonts
    unicode_fonts = regular not in {"Helvetica", "Helvetica-Bold"}
    base = ParagraphStyle("Base", fontName=regular, fontSize=9.4, leading=12.6,
                          textColor=(0.09, 0.12, 0.1), alignment=TA_LEFT, splitLongWords=1)
    name_style = ParagraphStyle("Name", parent=base, fontName=bold, fontSize=19, leading=22,
                                spaceAfter=1, textColor=(0.05, 0.08, 0.06))
    headline_style = ParagraphStyle("Headline", parent=base, fontName=italic, fontSize=10,
                                    leading=13, spaceAfter=2)
    contact_style = ParagraphStyle("Contact", parent=base, fontSize=8.6, leading=11.4,
                                   textColor=(0.28, 0.34, 0.29))
    section_style = ParagraphStyle("Section", parent=base, fontName=bold, fontSize=10.2,
                                   leading=13, spaceBefore=8, spaceAfter=3,
                                   textColor=(0.08, 0.18, 0.1))
    item_title_style = ParagraphStyle("ItemTitle", parent=base, fontName=bold, fontSize=9.7,
                                      leading=12.2, spaceBefore=4, spaceAfter=1)
    item_meta_style = ParagraphStyle("ItemMeta", parent=base, fontName=italic, fontSize=8.9,
                                     leading=11.4, textColor=(0.33, 0.4, 0.35), spaceAfter=1)
    bullet_style = ParagraphStyle("Bullet", parent=base, leftIndent=10, bulletIndent=1,
                                  spaceAfter=1.2)
    para_style = ParagraphStyle("Para", parent=base, fontSize=10, leading=14.2, spaceAfter=8)

    size = LETTER if page_size == "letter" else A4
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=size, leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=13 * mm, bottomMargin=15 * mm,
        title=f"{content.get('header', {}).get('name', '')} — "
              f"{'CV' if doc_type == 'cv' else 'Cover Letter'}",
        author=content.get("header", {}).get("name", ""),
        subject=(content.get("target") or {}).get("role", ""),
    )
    story = []
    header = content.get("header", {})
    story.append(_pdf_paragraph(header.get("name", ""), name_style, unicode_fonts))
    if header.get("headline"):
        story.append(_pdf_paragraph(header["headline"], headline_style, unicode_fonts))
    contact = "  |  ".join(bits for bits in (header.get("email"), header.get("phone"),
                                             header.get("location")) if bits)
    if contact:
        story.append(_pdf_paragraph(contact, contact_style, unicode_fonts))
    links = header.get("links") or []
    if links:
        story.append(_pdf_paragraph(
            "  |  ".join(f"{link.get('label')}: {link.get('url')}" for link in links),
            contact_style, unicode_fonts))

    if doc_type == "cover_letter":
        recipient = content.get("recipient", {})
        story.append(Spacer(1, 10))
        if content.get("date"):
            story.append(_pdf_paragraph(content["date"], para_style, unicode_fonts))
        if recipient.get("company"):
            story.append(_pdf_paragraph(recipient["company"], para_style, unicode_fonts))
        if recipient.get("role"):
            story.append(_pdf_paragraph(f"Re: {recipient['role']}", para_style, unicode_fonts))
        story.append(_pdf_paragraph(content.get("salutation", "Dear Hiring Manager,"),
                                    para_style, unicode_fonts))
        for paragraph_text in content.get("paragraphs", []):
            story.append(_pdf_paragraph(paragraph_text, para_style, unicode_fonts))
        story.append(Spacer(1, 4))
        story.append(_pdf_paragraph(content.get("closing", "Sincerely,"), para_style, unicode_fonts))
        story.append(_pdf_paragraph(content.get("signature", ""), item_title_style, unicode_fonts))
    else:
        if content.get("summary"):
            story.append(_pdf_paragraph("SUMMARY", section_style, unicode_fonts))
            story.append(_pdf_paragraph(content["summary"], base, unicode_fonts))
        for section in content.get("sections", []):
            story.append(_pdf_paragraph(str(section.get("title", "")).upper(), section_style,
                                        unicode_fonts))
            if section.get("type") == "skills":
                for group in section.get("groups", []):
                    group_text = (f"<b>{escape(_pdf_safe(group.get('category', 'Skills'), unicode_fonts))}:</b> "
                                  f"{escape(_pdf_safe(', '.join(group.get('items', [])), unicode_fonts))}")
                    story.append(Paragraph(group_text, base))
                continue
            for item in section.get("items", []):
                block = []
                title_text = escape(_pdf_safe(item.get("title", ""), unicode_fonts))
                if item.get("subtitle"):
                    title_text += " — " + escape(_pdf_safe(item["subtitle"], unicode_fonts))
                if item.get("dates"):
                    title_text += f"  ({escape(_pdf_safe(item['dates'], unicode_fonts))})"
                block.append(Paragraph(title_text, item_title_style))
                if item.get("location"):
                    block.append(_pdf_paragraph(item["location"], item_meta_style, unicode_fonts))
                story.append(KeepTogether(block))
                bullets = [str(b) for b in (item.get("bullets") or []) if str(b).strip()]
                if bullets:
                    items = [ListItem(_pdf_paragraph(bullet, bullet_style, unicode_fonts), leftIndent=8)
                             for bullet in bullets]
                    story.append(ListFlowable(items, bulletType="bullet", start="•",
                                              leftIndent=10, bulletFontSize=7, bulletOffsetY=-1))
                for link in item.get("links", []) or []:
                    story.append(_pdf_paragraph(f"Link: {link}", item_meta_style, unicode_fonts))
    try:
        doc.build(story, canvasmaker=_numbered_canvas_class())
    except Exception as exc:  # noqa: BLE001 - reportlab raises varied errors
        logger.warning("pdf_build_failed error=%s", type(exc).__name__)
        raise RenderError("The PDF could not be generated from this content.") from exc
    data = buffer.getvalue()
    ok, message = validate_pdf(data)
    if not ok:
        raise RenderError(message)
    return data


# --------------------------------------------------------------------------- LaTeX


def render_tex(content: dict, doc_type: str) -> bytes:
    def esc(value: str) -> str:
        text = strip_control_characters(str(value or ""))
        replacements = {
            "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#",
            "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\^{}",
        }
        for source, target in replacements.items():
            text = text.replace(source, target)
        return text

    header = content.get("header", {})
    lines = [
        r"\documentclass[10pt,a4paper]{article}",
        r"\usepackage[margin=1.6cm]{geometry}",
        r"\usepackage[T1]{fontenc}",
        r"\usepackage{lmodern}",
        r"\usepackage[utf8]{inputenc}",
        r"\usepackage{enumitem}",
        r"\usepackage[hidelinks]{hyperref}",
        r"\setlist[itemize]{leftmargin=1.2em,itemsep=1pt,topsep=1pt}",
        r"\pagestyle{empty}",
        r"\begin{document}",
        r"\begin{center}",
        r"{\LARGE \textbf{" + esc(header.get("name", "")) + r"}}\\[2pt]",
    ]
    if header.get("headline"):
        lines.append(r"{\large " + esc(header["headline"]) + r"}\\[2pt]")
    contact = " | ".join(bits for bits in (header.get("email"), header.get("phone"),
                                           header.get("location")) if bits)
    if contact:
        lines.append(esc(contact) + r"\\")
    for link in header.get("links", []) or []:
        lines.append(r"\href{" + esc(link.get("url", "")) + "}{" + esc(link.get("label", "")) + r"}\\")
    lines.append(r"\end{center}")

    if doc_type == "cover_letter":
        lines.extend([
            esc(content.get("date", "")) + r"\\[6pt]",
            esc((content.get("recipient") or {}).get("company", "")) + r"\\[6pt]",
            esc(content.get("salutation", "Dear Hiring Manager,")) + r"\\[6pt]",
        ])
        for paragraph_text in content.get("paragraphs", []):
            lines.extend([esc(paragraph_text) + r"\\[6pt]", ""])
        lines.extend([esc(content.get("closing", "Sincerely,")) + r"\\", esc(content.get("signature", ""))])
    else:
        if content.get("summary"):
            lines.extend([r"\section*{Summary}", esc(content["summary"])])
        for section in content.get("sections", []):
            lines.append(r"\section*{" + esc(section.get("title", "")) + "}")
            if section.get("type") == "skills":
                for group in section.get("groups", []):
                    lines.append(r"\textbf{" + esc(group.get("category", "Skills")) + "}: "
                                 + esc(", ".join(group.get("items", []))) + r"\\")
                continue
            for item in section.get("items", []):
                heading = esc(item.get("title", ""))
                if item.get("subtitle"):
                    heading += " --- " + esc(item["subtitle"])
                if item.get("dates"):
                    heading += " (" + esc(item["dates"]) + ")"
                lines.append(r"\textbf{" + heading + "}\\")
                bullets = [str(b) for b in (item.get("bullets") or []) if str(b).strip()]
                if bullets:
                    lines.append(r"\begin{itemize}")
                    for bullet in bullets:
                        lines.append(r"\item " + esc(bullet))
                    lines.append(r"\end{itemize}")
    lines.append(r"\end{document}")
    return "\n".join(lines).encode("utf-8")


def try_latex_pdf(tex_source: bytes, timeout: float = 20.0) -> bytes | None:
    """Compile LaTeX if a TeX installation is present; otherwise return None."""
    pdflatex = shutil.which("pdflatex")
    if not pdflatex:
        return None
    with tempfile.TemporaryDirectory(prefix="cp-tex-") as directory:
        source = Path(directory) / "document.tex"
        source.write_bytes(tex_source)
        try:
            subprocess.run(
                [pdflatex, "-interaction=nonstopmode", "-halt-on-error", "document.tex"],
                cwd=directory, timeout=timeout, check=False,
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            )
            output = Path(directory) / "document.pdf"
            if output.exists():
                data = output.read_bytes()
                ok, _ = validate_pdf(data)
                return data if ok else None
        except (subprocess.SubprocessError, OSError):
            logger.info("latex_compile_failed")
        return None


# --------------------------------------------------------------------------- validation


def validate_pdf(data: bytes) -> tuple[bool, str]:
    if not data or not data.startswith(b"%PDF-"):
        return False, "Generated PDF is missing a valid header."
    if len(data) < 800:
        return False, "Generated PDF is unexpectedly small."
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) < 1:
            return False, "Generated PDF has no pages."
        text = "".join((page.extract_text() or "") for page in reader.pages[:2])
        if len(text.strip()) < 20:
            return False, "Generated PDF does not contain extractable text."
        return True, f"PDF valid with {len(reader.pages)} page(s)."
    except Exception as exc:  # noqa: BLE001
        logger.warning("pdf_validation_failed error=%s", type(exc).__name__)
        return False, "Generated PDF could not be opened for validation."


def validate_docx(data: bytes) -> tuple[bool, str]:
    if not data or not data.startswith(b"PK"):
        return False, "Generated DOCX is missing a valid archive header."
    try:
        import zipfile

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if "word/document.xml" not in archive.namelist():
                return False, "Generated DOCX does not contain a document body."
        import docx

        document = docx.Document(io.BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        if len(text.strip()) < 20:
            return False, "Generated DOCX does not contain readable text."
        return True, "DOCX valid and readable."
    except Exception as exc:  # noqa: BLE001
        logger.warning("docx_validation_failed error=%s", type(exc).__name__)
        return False, "Generated DOCX could not be opened for validation."


def render_document(content: dict, doc_type: str, fmt: str) -> tuple[bytes, str]:
    """Render and validate one export format. Returns (bytes, validation message)."""
    if fmt == "docx":
        data = render_docx(content, doc_type)
        ok, message = validate_docx(data)
        if not ok:
            raise RenderError(message)
        return data, message
    if fmt == "txt":
        data = render_text(content, doc_type)
        return data, f"Plain text generated ({len(data)} bytes)."
    if fmt == "tex":
        return render_tex(content, doc_type), "LaTeX source generated."
    if fmt == "pdf":
        data = render_pdf(content, doc_type)
        ok, message = validate_pdf(data)
        if not ok:
            raise RenderError(message)
        return data, message
    raise RenderError(f"Unsupported format: {fmt}")
